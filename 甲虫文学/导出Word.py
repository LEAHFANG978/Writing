"""从当前Markdown生成Word；不改小说正文，不依赖网络。"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent


def plain(text):
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1（\2）', text)
    return text.replace('**', '').replace('`', '')


def export(source, target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.6)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for level in range(1, 5):
        doc.styles[f'Heading {level}'].element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    lines = source.read_text(encoding='utf-8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [value.strip() for value in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', value) for value in cells):
                    rows.append(cells)
                index += 1
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Table Grid'
            for number, cells in enumerate(rows):
                row = table.add_row()
                for column, value in enumerate(cells):
                    row.cells[column].text = plain(value)
                    for paragraph in row.cells[column].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(7.5)
                            run.bold = number == 0
                if number == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            continue
        heading = re.match(r'^(#{1,4})\s+(.*)', line)
        if heading:
            doc.add_heading(plain(heading[2]), level=len(heading[1]))
        elif line.strip() == '---':
            doc.add_page_break()
        elif line.strip() and not line.startswith('```'):
            doc.add_paragraph(plain(line))
        index += 1
    doc.save(target)
    restored = Document(target)
    assert len(restored.paragraphs) > 10
    print(f'已生成并重新解析：{target.name}')


if __name__ == '__main__':
    master = ROOT / '甲虫文学｜快消职场荒诞小说创作圣经.md'
    export(master, master.with_suffix('.docx'))
    for source in sorted((ROOT / '小说正文').glob('*.md')):
        target = source.with_suffix('.docx')
        if target.exists():
            print(f'保留既有小说Word，不覆盖：{target.name}')
            continue
        export(source, target)
